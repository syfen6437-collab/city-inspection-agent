from __future__ import annotations

import hashlib
import io
import os
import re
import shutil
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable

from docx import Document


class UnsupportedDocumentError(RuntimeError):
    pass


@dataclass
class DocumentBlock:
    source_id: str
    text: str
    kind: str = "paragraph"
    table_index: int | None = None
    row_index: int | None = None
    page: int | None = None

    def to_dict(self) -> dict:
        return {key: value for key, value in asdict(self).items() if value is not None}


@dataclass
class ParsedReport:
    report_id: str
    file_name: str
    extension: str
    blocks: list[DocumentBlock]
    source_metadata: dict[str, str | int | None]

    @property
    def raw_text(self) -> str:
        return "\n".join(block.text for block in self.blocks if block.text)

    def to_dict(self) -> dict:
        return {
            "report_id": self.report_id,
            "file_name": self.file_name,
            "extension": self.extension,
            "source_metadata": self.source_metadata,
            "blocks": [block.to_dict() for block in self.blocks],
        }


def normalize_text(value: object) -> str:
    text = str(value or "").replace("\u00a0", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _block_id(kind: str, index: int, row: int | None = None) -> str:
    return f"{kind}:{index}" if row is None else f"{kind}:{index}:{row}"


def parse_docx_bytes(data: bytes, file_name: str) -> ParsedReport:
    document = Document(io.BytesIO(data))
    blocks: list[DocumentBlock] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = normalize_text(paragraph.text)
        if text:
            blocks.append(DocumentBlock(_block_id("paragraph", index), text, "paragraph"))
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            values = [normalize_text(cell.text) for cell in row.cells]
            text = " | ".join(values).strip(" |")
            if text:
                blocks.append(DocumentBlock(_block_id("table", table_index, row_index), text, "table", table_index, row_index))
    digest = hashlib.sha256(data).hexdigest()[:16]
    return ParsedReport(digest, Path(file_name).name, ".docx", blocks, extract_source_metadata(blocks))


def _convert_doc_with_word(data: bytes, file_name: str) -> bytes:
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise UnsupportedDocumentError("解析 .doc 需要安装 pywin32，且不能读取原始二进制文档") from exc
    temp_dir = Path(tempfile.mkdtemp(prefix="city-agent-doc-"))
    source = temp_dir / Path(file_name).name
    target = temp_dir / (source.stem + ".docx")
    source.write_bytes(data)
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(source), ReadOnly=True, AddToRecentFiles=False)
        document.SaveAs2(str(target), FileFormat=16)
        document.Close(False)
        document = None
        return target.read_bytes()
    except Exception as exc:
        raise UnsupportedDocumentError(f"Word 无法转换 .doc 文件：{exc}") from exc
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit(False)
            except Exception:
                pass
        shutil.rmtree(temp_dir, ignore_errors=True)


def parse_document_bytes(data: bytes, file_name: str) -> ParsedReport:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".docx":
        return parse_docx_bytes(data, file_name)
    if suffix == ".doc":
        converted = _convert_doc_with_word(data, file_name)
        report = parse_docx_bytes(converted, file_name)
        report.extension = ".doc"
        return report
    raise UnsupportedDocumentError(f"不支持的文件类型：{suffix or '<无扩展名>'}")


def parse_document_path(path: str | Path) -> ParsedReport:
    source = Path(path)
    return parse_document_bytes(source.read_bytes(), source.name)


def iter_document_paths(root: str | Path) -> Iterable[Path]:
    base = Path(root)
    for path in sorted(base.rglob("*")):
        if path.is_file() and path.suffix.lower() in {".docx", ".doc"}:
            yield path


def extract_source_metadata(blocks: list[DocumentBlock]) -> dict[str, str | int | None]:
    """Extract non-sensitive routing hints only; prediction fields still come from the model."""
    text = "\n".join(block.text for block in blocks[:80])
    date_match = re.search(r"(20\d{2})\s*[年/-]\s*(\d{1,2})\s*[月/-]\s*(\d{1,2})\s*日?", text)
    report_match = re.search(r"报告编号\s*[:：]\s*([A-Za-z0-9-]+)", text)
    bci_match = re.search(r"\bBCI\s*[=:：]\s*([0-9]+(?:\.[0-9]+)?)", text, re.IGNORECASE)
    return {
        "inspection_year": int(date_match.group(1)) if date_match else None,
        "inspection_date_hint": "-".join(date_match.groups()) if date_match else None,
        "report_number_hint": report_match.group(1) if report_match else None,
        "bci_hint": float(bci_match.group(1)) if bci_match else None,
    }

