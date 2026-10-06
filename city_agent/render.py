from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .adapter import redact_text
from .schema import PredictionRecord


def _safe(value: object) -> str:
    if value is None:
        return "未提取"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, dict):
        return "；".join(f"{key}：{_safe(item)}" for key, item in value.items()) or "未提取"
    if isinstance(value, (list, tuple)):
        return "；".join(_safe(item) for item in value) or "未提取"
    return redact_text(str(value))


def _remove_paragraph_borders(element: object) -> None:
    """Remove Word's built-in title rule from a style or paragraph."""
    ppr = element.get_or_add_pPr()
    for child in list(ppr):
        if child.tag == qn("w:pBdr"):
            ppr.remove(child)


def _table(document: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EEF5")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = _safe(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
            for row in table.rows:
                row.cells[index].width = Inches(width)


def render_prediction_docx(record: PredictionRecord, output_path: str | Path) -> None:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    title_style = document.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.size = Pt(16)
    _remove_paragraph_borders(title_style._element)
    heading_style = document.styles["Heading 1"]
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _remove_paragraph_borders(title._p)
    run = title.add_run("城市基础设施定检报告信息提取结果")
    run.bold = True
    document.add_paragraph(f"源文件：{_safe(record.file_name)}")
    document.add_paragraph(f"处理状态：{_safe(record.status)}")
    if record.error:
        document.add_paragraph(f"错误信息：{_safe(record.error)}")
    if record.prediction is None:
        document.add_paragraph("模型未生成预测结果。本文件仅保留处理状态，不代表评测答案。")
    else:
        prediction = record.prediction
        document.add_heading("一、概要字段", level=1)
        _table(
            document,
            ["字段", "内容"],
            [
                ["桥梁/设施名称", prediction.bridge_name],
                ["报告编号", prediction.report_number],
                ["检测日期", prediction.inspection_date],
                ["检测年份", prediction.inspection_year],
                ["总体评分", prediction.overall_score],
                ["总体等级", prediction.overall_grade],
                ["构件评分", prediction.component_scores],
            ], widths=[1.6, 4.9],
        )
        document.add_heading("二、摘要与重点风险", level=1)
        if prediction.summary:
            document.add_paragraph(_safe(prediction.summary))
        for item in prediction.key_risks:
            document.add_paragraph(_safe(item), style="List Bullet")
        if not prediction.summary and not prediction.key_risks:
            document.add_paragraph("未提取")
        document.add_heading("三、病害明细", level=1)
        _table(
            document,
            ["位置", "类型", "描述", "是否新增", "前次状态", "发展程度", "测量", "证据"],
            [
                [
                    item.location,
                    item.disease_type,
                    item.description,
                    item.is_new,
                    item.previous_status,
                    item.development,
                    item.measurement,
                    ", ".join(item.evidence_ids),
                ]
                for item in prediction.diseases
            ],
            widths=[1.05, 0.72, 1.05, 0.55, 0.82, 0.72, 0.8, 0.79],
        )
        document.add_heading("四、维修建议与依据", level=1)
        for item in prediction.recommendations:
            document.add_paragraph(_safe(item), style="List Bullet")
        for item in prediction.standards:
            document.add_paragraph(_safe(item), style="List Bullet 2")
        if not prediction.recommendations and not prediction.standards:
            document.add_paragraph("未提取")
        document.add_heading("五、证据链", level=1)
        evidence_rows = []
        seen_evidence: set[str] = set()
        for item in prediction.evidence:
            if item.source_id in seen_evidence:
                continue
            seen_evidence.add(item.source_id)
            evidence_rows.append([item.source_id, item.kind, item.quote])
        _table(
            document,
            ["来源编号", "类型", "原文摘录"],
            evidence_rows or [["未提取", "", ""]],
            widths=[1.15, 0.85, 4.5],
        )
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
