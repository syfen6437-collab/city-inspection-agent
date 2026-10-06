from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from city_agent.adapter import redact_public, validate_predictions_v1, write_predictions_v1
from city_agent.document_parser import parse_document_bytes
from city_agent.extractor import extract_report
from city_agent.index import NgramIndex
from city_agent.schema import DiseaseRecord, PredictionRecord, ReportPrediction


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("桥梁名称：测试桥")
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ("位置", "病害类型", "描述", "状态")):
        cell.text = value
    row = table.add_row().cells
    for cell, value in zip(row, ("1号桥墩", "裂缝", "竖向裂缝 0.20mm", "新增")):
        cell.text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class CoreTests(unittest.TestCase):
    def test_extractor_merges_chunk_outputs_and_filters_evidence(self) -> None:
        report = parse_document_bytes(make_docx(), "sample.docx")

        class FakeModel:
            metadata = {"model_name": "fake"}

            def __init__(self):
                self.calls = 0

            def generate_json(self, _system, _user):
                self.calls += 1
                if self.calls == 1:
                    return ({"bridge_name": "测试桥", "overall_grade": "B"}, {"latency_ms": 1})
                return ({
                    "diseases": [{"location": "1号桥墩", "disease_type": "裂缝", "evidence_ids": ["table:0:1", "bad"]}],
                    "recommendations": ["观察"],
                    "evidence": [
                        {"source_id": "table:0:1", "quote": "模型可能改写的内容"},
                        {"source_id": "table:0:1", "quote": "重复来源"},
                        {"source_id": "bad", "quote": "无效"},
                    ],
                }, {"latency_ms": 2})

        model = FakeModel()
        record = extract_report(report, model)
        self.assertEqual(record.status, "success")
        self.assertEqual(model.calls, 2)
        self.assertEqual(record.prediction.bridge_name, "测试桥")
        self.assertEqual(record.prediction.diseases[0].evidence_ids, ["table:0:1"])
        self.assertEqual([item.source_id for item in record.prediction.evidence], ["table:0:1"])

    def test_docx_parser_and_source_ids(self) -> None:
        report = parse_document_bytes(make_docx(), "sample.docx")
        self.assertEqual(report.extension, ".docx")
        self.assertTrue(any(block.source_id == "paragraph:0" for block in report.blocks))
        self.assertTrue(any(block.source_id == "table:0:1" for block in report.blocks))
        self.assertIn("裂缝", report.raw_text)

    def test_index_returns_source_evidence(self) -> None:
        report = parse_document_bytes(make_docx(), "sample.docx")
        index = NgramIndex.from_reports([report])
        hits = index.search("竖向裂缝", limit=2)
        self.assertTrue(hits)
        self.assertIn("裂缝", hits[0].quote)
        self.assertEqual(index.search("不存在的查询", report_id=report.report_id), [])

    def test_schema_round_trip(self) -> None:
        value = ReportPrediction.from_dict({
            "bridge_name": "测试桥",
            "overall_score": "87.5",
            "inspection_year": "2024",
            "diseases": [{"location": "桥墩", "disease_type": "裂缝", "evidence_ids": ["table:0:1"]}],
        })
        self.assertEqual(value.overall_score, 87.5)
        self.assertEqual(value.inspection_year, 2024)
        self.assertIsInstance(value.diseases[0], DiseaseRecord)
        with self.assertRaisesRegex(ValueError, "evidence_ids 必须是数组"):
            ReportPrediction.from_dict({"diseases": [{"evidence_ids": "paragraph:1"}]})

    def test_public_redaction_and_validation(self) -> None:
        public = redact_public({"phone": "13812345678", "nested": "联系人：张三，联系地址：重庆市渝中区某路 12 号"})
        self.assertEqual(public["phone"], "[已脱敏]")
        self.assertNotIn("张三", json.dumps(public, ensure_ascii=False))
        self.assertNotIn("重庆市渝中区某路", json.dumps(public, ensure_ascii=False))
        self.assertNotIn("13812345678", json.dumps(public, ensure_ascii=False))
        record = PredictionRecord("sample.docx", "rid", "blocked", error="model unavailable")
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "predictions.v1.json"
            write_predictions_v1([record], path, "dataset.zip", "test", {"model_name": "local"})
            self.assertEqual(validate_predictions_v1(path), [])


if __name__ == "__main__":
    unittest.main()
