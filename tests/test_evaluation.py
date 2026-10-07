from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document

from city_agent.document_parser import DocumentBlock, ParsedReport, parse_document_bytes
from city_agent.evaluation import _field_match, evaluate_predictions, label_consistency_issues, parse_label
from city_agent.training_data import build_group_split, pair_training_files
from city_agent.schema import ReportPrediction


class EvaluationTests(unittest.TestCase):
    def test_extended_official_schema_round_trip(self):
        payload = {"bridge_name": "测试桥", "previous_overall_score": "87.00", "previous_overall_grade": "B级", "disease_trend": "原文趋势", "detailed_conclusion": ["原文结论"], "disease_causes": ["原文成因"], "disposal_recommendations": ["原文建议"], "safety_impacts": ["原文影响"], "recommendation_details": [{"category": "尽快维修", "content": "维修", "location": "桥面"}]}
        prediction = ReportPrediction.from_dict(payload)
        for key, value in payload.items():
            self.assertEqual(prediction.to_dict()[key], value)
        with self.assertRaises(ValueError):
            ReportPrediction.from_dict({"recommendation_details": [{"content": 123}]})

    def test_label_conflicts_are_flagged_without_rewriting(self):
        label = {"summary": {"总体评分": "90.75", "上一次总体评分": "74.34"}, "detailed_conclusion": ["总体技术状况评分90.75分。与上一年度（2012年，总体评分90.78分）相比。"]}
        issues = label_consistency_issues(label)
        self.assertEqual([item["field"] for item in issues], ["上一次总体评分"])
        self.assertEqual(label["summary"]["上一次总体评分"], "74.34")

    def test_evaluation_end_to_end_and_test_split_rejected(self):
        label_doc = Document()
        table = label_doc.add_table(rows=2, cols=3)
        for cell, text in zip(table.rows[0].cells, ["字段", "示例", "说明"]):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, ["总体评分", "86.43", "score"]):
            cell.text = text
        label_doc.add_paragraph("（1）详细结论")
        label_doc.add_paragraph("a conclusion")
        label_doc.add_paragraph("（2）建议明细")
        label_data = io.BytesIO()
        label_doc.save(label_data)
        source_doc = Document()
        source_doc.add_paragraph("source")
        source_data = io.BytesIO()
        source_doc.save(source_data)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "official.zip"
            with zipfile.ZipFile(archive, "w") as handle:
                handle.writestr("root/赛题一_训练集/2013年/测试桥.docx", source_data.getvalue())
                handle.writestr("root/赛题一_训练集标签/2013年/测试桥-无对比年度的信息提取报告.docx", label_data.getvalue())
            prediction = root / "predictions.json"
            payload = {"split": "train", "records": [{"file_name": "测试桥.docx", "status": "success", "prediction": {"overall_score": 86.43, "detailed_conclusion": ["a conclusion"]}}]}
            prediction.write_text(json.dumps(payload), encoding="utf-8")
            result = evaluate_predictions(archive, prediction, root)
            self.assertTrue(result["complete"])
            self.assertEqual(result["matched"], 1)
            self.assertEqual(result["summary_field_accuracy"], 1)
            self.assertEqual(result["detailed_character_f1"], 1)
            payload["records"].append({"file_name": "not_found.docx", "status": "failed"})
            prediction.write_text(json.dumps(payload), encoding="utf-8")
            result = evaluate_predictions(archive, prediction, root)
            self.assertFalse(result["complete"])
            self.assertEqual(result["evaluation_coverage"], 0.5)
            payload["split"] = "test"
            prediction.write_text(json.dumps(payload), encoding="utf-8")
            with patch("city_agent.evaluation.zipfile.ZipFile", side_effect=AssertionError("must not open labels")):
                with self.assertRaisesRegex(ValueError, "split=train"):
                    evaluate_predictions(archive, prediction, root)

    def test_fields_are_not_matched_against_unrelated_values(self):
        result = _field_match({"prediction": {"bridge_name": "B级", "overall_grade": "A级", "overall_score": 86.4}}, {"summary": {"总体等级": "B级", "总体评分": "86.43"}})
        self.assertEqual(result["summary_field_accuracy"], 0)

    def test_nested_rows_one_to_one_and_evidence(self):
        disease = {"location": "1#墩", "disease_type": "裂缝", "description": "竖向裂缝", "evidence_ids": ["table:0:1"]}
        record = {"prediction": {"diseases": [disease], "evidence": [{"source_id": "table:0:1", "quote": "竖向裂缝"}]}}
        label = {"summary": {}, "diseases": [disease]}
        report = ParsedReport("rid", "x.docx", ".docx", [DocumentBlock("table:0:1", "1#墩竖向裂缝")], {})
        result = _field_match(record, label, report)
        self.assertEqual(result["disease_record_recall"], 1)
        self.assertEqual(result["evidence_coverage"], 1)
        label["diseases"].append(disease)
        self.assertEqual(_field_match(record, label, report)["disease_record_recall"], 0.5)
        record["prediction"]["evidence"][0]["quote"] = "不存在的原文"
        result = _field_match(record, label, report)
        self.assertEqual(result["evidence_coverage"], 0)
        self.assertEqual(result["invalid_evidence_count"], 1)

    def test_summary_dates_grades_and_missing_values(self):
        record = {"prediction": {"inspection_date": "2013年10月", "overall_grade": "B", "overall_score": 86.43}}
        label = {"summary": {"报告日期": "2013年10月", "总体等级": "B级", "总体评分": "86.430", "上一次总体评分": "无"}}
        self.assertEqual(_field_match(record, label)["summary_field_accuracy"], 0.75)
        record["prediction"]["inspection_date"] = "2013年10月01日"
        self.assertFalse(_field_match(record, label)["fields"]["报告日期"])

    def test_parser_preserves_cell_positions_and_body_order(self):
        doc = Document()
        doc.add_paragraph("before")
        table = doc.add_table(rows=1, cols=4)
        for cell, text in zip(table.rows[0].cells, ["序号", "病害部位", "病害类型", "病害描述"]):
            cell.text = text
        row = table.add_row().cells
        for cell, text in zip(row, ["1", "", "裂缝", "desc | containing separator"]):
            cell.text = text
        doc.add_paragraph("after")
        buffer = io.BytesIO()
        doc.save(buffer)
        report = parse_document_bytes(buffer.getvalue(), "x.docx")
        self.assertEqual([block.kind for block in report.blocks], ["paragraph", "table", "table", "paragraph"])
        parsed = parse_label(report)
        self.assertEqual(parsed["diseases"][0]["location"], "")
        self.assertEqual(parsed["diseases"][0]["description"], "desc | containing separator")

    def test_label_table_headers_not_row_keywords(self):
        blocks = [
            DocumentBlock("table:0:0", "字段 | 示例 | 说明", "table", 0, 0),
            DocumentBlock("table:0:1", "报告日期 | 2014年11月 | date", "table", 0, 1),
            DocumentBlock("table:0:2", "上部结构等级 | C级 | grade", "table", 0, 2),
            DocumentBlock("table:1:0", "序号 | 建议类别 | 建议内容 | 病害部位", "table", 1, 0),
            DocumentBlock("table:1:1", "1 | 立即处置 | 灌浆 | 梁底", "table", 1, 1),
        ]
        parsed = parse_label(ParsedReport("rid", "x.docx", ".docx", blocks, {}))
        self.assertEqual(len(parsed["summary"]), 2)
        self.assertEqual(parsed["recommendations"][0]["category"], "立即处置")
        blocks.pop(0)
        parsed = parse_label(ParsedReport("rid", "x.docx", ".docx", blocks, {}))
        self.assertEqual(len(parsed["summary"]), 2)
        self.assertEqual(parsed["warnings"], [])

    def test_recommendation_counts_are_compared_as_counts(self):
        record = {"prediction": {"recommendation_details": [{"category": "尽快维修"}]}}
        label = {"summary": {"建议": "0条立即处置、1条尽快维修、0条预防性养护建议"}}
        self.assertEqual(_field_match(record, label)["summary_field_accuracy"], 1)
        label["summary"]["建议"] = "1条立即处置、1条尽快维修"
        self.assertEqual(_field_match(record, label)["summary_field_accuracy"], 0)


class TrainingInventoryTests(unittest.TestCase):
    @staticmethod
    def source(year, name):
        return f"root/赛题一_训练集/{year}年/{name}.doc"

    @staticmethod
    def label(year, name):
        return f"root/赛题一_训练集标签/{year}年/{name}-含有对比年度的信息提取报告.docx"

    def test_same_year_pairing_and_nested_docx_suffix(self):
        originals = [self.source(2013, "12-048土主互通立交报告"), self.source(2014, "12-048土主互通立交报告")]
        labels = [self.label(2013, "12-048土主互通立交.docx"), self.label(2014, "12-048土主互通立交.docx")]
        result = pair_training_files(originals, labels)
        self.assertEqual(result["paired_count"], 2)
        self.assertEqual(result["pairs"][0]["year"], "2013年")

    def test_many_to_one_or_ambiguous_labels_not_forced(self):
        originals = [self.source(2013, "测试桥"), self.source(2013, "测试桥报告")]
        labels = [self.label(2013, "测试桥")]
        self.assertEqual(pair_training_files(originals, labels)["paired_count"], 0)
        self.assertEqual(pair_training_files(originals[:1], labels + [self.label(2013, "测试桥报告")])["paired_count"], 0)

    def test_different_chainages_not_paired(self):
        result = pair_training_files([self.source(2012, "成渝K349+318小桥")], [self.label(2012, "成渝K349+319")])
        self.assertEqual(result["paired_count"], 0)

    def test_connected_bridge_groups_stable_and_disjoint(self):
        pairs = [
            {"train_file": self.source(2012, "青果林中桥"), "label_file": self.label(2012, "青果林中桥")},
            {"train_file": self.source(2014, "100-青果林中桥(K8+150)2014.11.21"), "label_file": self.label(2014, "100-青果林中桥(K8+150)")},
            {"train_file": self.source(2014, "101-别的桥"), "label_file": self.label(2014, "101-别的桥")},
        ]
        split = build_group_split(pairs)
        self.assertEqual(split["group_count"], 2)
        train_groups = {item["bridge_group"] for item in split["train"]}
        validation_groups = {item["bridge_group"] for item in split["validation"]}
        self.assertFalse(train_groups & validation_groups)
        reverse = build_group_split(list(reversed(pairs)))
        self.assertEqual(validation_groups, {item["bridge_group"] for item in reverse["validation"]})


if __name__ == "__main__":
    unittest.main()
